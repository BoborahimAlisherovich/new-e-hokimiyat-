#!/usr/bin/env python3
"""
OneID texnologik yorqnomasini o'qish uchun skript
"""

import sys
from docx import Document

def read_oneid_document():
    try:
        doc = Document('TexnologikYoriqnoma.docx')
        
        print("=" * 80)
        print("ONEID TEXNOLOGIK YORQNOMASI MA'LUMOTLARI")
        print("=" * 80)
        
        for i, paragraph in enumerate(doc.paragraphs):
            text = paragraph.text.strip()
            if text:
                print(f"{i+1:3d}: {text}")
        
        print("\n" + "=" * 80)
        print("JADVAL MA'LUMOTLARI")
        print("=" * 80)
        
        for table_idx, table in enumerate(doc.tables):
            print(f"\nJadval {table_idx + 1}:")
            for row_idx, row in enumerate(table.rows):
                row_data = [cell.text.strip() for cell in row.cells]
                if any(row_data):  # Faqat bo'sh bo'lmagan qatorlarni ko'rsatish
                    print(f"  {row_idx + 1}: {' | '.join(row_data)}")
        
    except Exception as e:
        print(f"Xatolik yuz berdi: {e}")
        return False
    
    return True

if __name__ == "__main__":
    read_oneid_document()
