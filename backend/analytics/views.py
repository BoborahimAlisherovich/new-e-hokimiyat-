"""
Analytics views for E-Hokimiyat API.
"""

from rest_framework import views, status
from rest_framework.response import Response
from rest_framework.permissions import IsAuthenticated
from django.http import HttpResponse
from django.db.models import Count, Q, Avg, F
from django.db.models.functions import TruncDate
from django.utils import timezone
from datetime import timedelta
from io import BytesIO

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

from tasks.models import Task, TaskOrganization
from organizations.models import Organization
from users.models import User
from telegram_bot.models import BotRegion, TelegramAppeal

def _organization_stats():
    """
    Tashkilotlar kesimidagi ko'rsatkichlar — IKKI so'rov (ilgari 4N+1).

    Ilgari har bir tashkilot uchun to'rtta alohida `.count()` bajarilardi:
    `for org in organizations:` ichida total/completed/overdue/in_progress.
    200 tashkilotda bu 800+ so'rov degani. Xuddi shu kod eksport yo'lida
    ham so'zma-so'z takrorlangan edi.
    """
    grouped = {
        row['organization_id']: row
        for row in (
            TaskOrganization.objects
            .filter(organization__is_active=True)
            .values('organization_id')
            .annotate(
                total=Count('id'),
                completed=Count('id', filter=Q(status__in=['BAJARILDI', 'NAZORATDAN_YECHILDI'])),
                overdue=Count('id', filter=Q(status='MUDDATI_KECH')),
                in_progress=Count('id', filter=Q(status='IJRODA')),
            )
        )
    }

    stats = []
    # Topshiriqsiz tashkilotlar ham ro'yxatda qoladi (eski xatti-harakat).
    for org in Organization.objects.filter(is_active=True).only('id', 'name', 'short_name'):
        row = grouped.get(org.id, {})
        total = row.get('total', 0) or 0
        completed = row.get('completed', 0) or 0
        overdue = row.get('overdue', 0) or 0
        in_progress = row.get('in_progress', 0) or 0

        completion_rate = (completed / total * 100) if total else 0
        penalty = (overdue / total * 50) if total else 0
        rating = max(0.0, completion_rate - penalty)

        stats.append({
            'id': str(org.id),
            'name': org.name,
            'short_name': org.short_name,
            'total': total,
            'completed': completed,
            'in_progress': in_progress,
            'overdue': overdue,
            'completion_rate': round(completion_rate, 1),
            'rating': round(rating, 1),
        })

    stats.sort(key=lambda x: x['rating'], reverse=True)
    return stats




class IsHokimOrHokimlikMasul:
    """Permission for viewing analytics."""
    def has_permission(self, request, view):
        # Barcha autentifikatsiya qilingan foydalanuvchilar dashboard analytics ko'ra oladi
        if not request.user or not request.user.is_authenticated:
            return False
        # Lekin to'liq ma'lumot faqat yuqori rollarga
        return request.user.role in ['HOKIM', 'HOKIMLIK_MASUL', 'ADMIN', 'SUPERADMIN', 'SECTOR_LEADER', 'ORGANIZATION_HEAD', 'TASHKILOT_RAHBARI', 'TASHKILOT_MASUL']


class DashboardAnalyticsView(views.APIView):
    """
    Dashboard analytics endpoint.
    
    GET /api/analytics/dashboard/
    
    Tashkilot rahbari/mas'uli uchun faqat o'z tashkilotiga oid statistikani qaytaradi.
    """
    permission_classes = [IsAuthenticated, IsHokimOrHokimlikMasul]
    
    def get(self, request):
        today = timezone.now().date()
        thirty_days_ago = today - timedelta(days=30)
        user = request.user
        
        is_org_user = user.role in ['TASHKILOT_RAHBARI', 'TASHKILOT_MASUL'] and user.organization
        
        if is_org_user:
            # Tashkilot xodimlari uchun TaskOrganization orqali hisoblash
            task_orgs = TaskOrganization.objects.filter(organization=user.organization)
            
            total_tasks = task_orgs.count()
            new_tasks = task_orgs.filter(status='YANGI').count()
            in_progress_tasks = task_orgs.filter(status='IJRODA').count()
            completed_tasks = task_orgs.filter(status__in=['BAJARILDI', 'NAZORATDAN_YECHILDI']).count()
            pending_approval = task_orgs.filter(status='TEKSHIRUVDA').count()
            resubmitted = task_orgs.filter(status='QAYTA_IJROGA_YUBORILDI').count()
            overdue_tasks = task_orgs.filter(status='MUDDATI_KECH').count()
            
            created_recently = task_orgs.filter(
                assigned_at__date__gte=thirty_days_ago
            ).count() if hasattr(TaskOrganization, 'assigned_at') else task_orgs.filter(
                task__created_at__date__gte=thirty_days_ago
            ).count()
            
            tasks_by_priority = task_orgs.values('task__priority').annotate(
                count=Count('id')
            ).values_list('task__priority', 'count').order_by()
            tasks_by_priority = [{'priority': p, 'count': c} for p, c in tasks_by_priority]
            
            total_users = user.organization.employees.filter(status='FAOL').count()
            active_users = total_users
            total_organizations = 1
            
            tasks_trend = task_orgs.filter(
                task__created_at__date__gte=thirty_days_ago
            ).annotate(
                date=TruncDate('task__created_at')
            ).values('date').annotate(
                created=Count('id')
            ).order_by('date')
            
            completed_trend = task_orgs.filter(
                status__in=['BAJARILDI', 'NAZORATDAN_YECHILDI'],
                completed_at__date__gte=thirty_days_ago
            ).annotate(
                date=TruncDate('completed_at')
            ).values('date').annotate(
                completed=Count('id')
            ).order_by('date')
        else:
            # Admin rollar uchun global statistika
            total_tasks = Task.objects.count()
            new_tasks = Task.objects.filter(status='YANGI').count()
            in_progress_tasks = Task.objects.filter(status='IJRODA').count()
            completed_tasks = Task.objects.filter(status__in=['BAJARILDI', 'NAZORATDAN_YECHILDI']).count()
            pending_approval = Task.objects.filter(status='YANGI').count()
            resubmitted = Task.objects.filter(status='QAYTA_IJROGA_YUBORILDI').count()
            overdue_tasks = Task.objects.filter(status='MUDDATI_KECH').count()
            
            created_recently = Task.objects.filter(
                created_at__date__gte=thirty_days_ago
            ).count()
            
            tasks_by_priority = list(Task.objects.values('priority').annotate(count=Count('id')))
            
            total_users = User.objects.count()
            active_users = User.objects.filter(status='FAOL').count()
            total_organizations = Organization.objects.filter(is_active=True).count()
            
            tasks_trend = Task.objects.filter(
                created_at__date__gte=thirty_days_ago
            ).annotate(
                date=TruncDate('created_at')
            ).values('date').annotate(
                created=Count('id')
            ).order_by('date')
            
            completed_trend = Task.objects.filter(
                closed_at__date__gte=thirty_days_ago
            ).annotate(
                date=TruncDate('closed_at')
            ).values('date').annotate(
                completed=Count('id')
            ).order_by('date')
        
        return Response({
            'summary': {
                'total_tasks': total_tasks,
                'new_tasks': new_tasks,
                'in_progress_tasks': in_progress_tasks,
                'completed_tasks': completed_tasks,
                'created_recently': created_recently,
                'pending_approval': pending_approval,
                'resubmitted': resubmitted,
                'overdue_tasks': overdue_tasks,
                'total_users': total_users,
                'active_users': active_users,
                'total_organizations': total_organizations,
            },
            'tasks_by_priority': list(tasks_by_priority) if not isinstance(tasks_by_priority, list) else tasks_by_priority,
            'tasks_trend': list(tasks_trend),
            'completed_trend': list(completed_trend),
        })


class OrgDashboardView(views.APIView):
    """
    Tashkilot rahbari/mas'uli uchun maxsus dashboard endpoint.
    
    GET /api/analytics/org-dashboard/
    
    O'z tashkilotiga tegishli topshiriqlar, murojaatlar va xodimlar statistikasini qaytaradi.
    """
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        user = request.user
        org = user.organization
        
        if not org:
            return Response(
                {'detail': "Sizning tashkilotingiz topilmadi"},
                status=status.HTTP_400_BAD_REQUEST
            )
        
        today = timezone.now().date()
        thirty_days_ago = today - timedelta(days=30)
        
        # ---- Task statistikasi (TaskOrganization orqali) ----
        task_orgs = TaskOrganization.objects.filter(organization=org).select_related('task', 'assigned_to')
        
        total_tasks = task_orgs.count()
        new_tasks = task_orgs.filter(status='YANGI').count()
        in_progress = task_orgs.filter(status='IJRODA').count()
        in_review = task_orgs.filter(status='TEKSHIRUVDA').count()
        completed = task_orgs.filter(status__in=['BAJARILDI', 'NAZORATDAN_YECHILDI']).count()
        overdue = task_orgs.filter(status='MUDDATI_KECH').count()
        resubmitted = task_orgs.filter(status='QAYTA_IJROGA_YUBORILDI').count()
        
        completion_rate = round((completed / total_tasks * 100), 1) if total_tasks > 0 else 0
        
        # Oxirgi 5 ta faol topshiriq
        recent_tasks = task_orgs.filter(
            status__in=['YANGI', 'IJRODA', 'QAYTA_IJROGA_YUBORILDI', 'MUDDATI_KECH', 'TEKSHIRUVDA']
        ).order_by('task__deadline')[:5]
        
        recent_tasks_data = []
        for to in recent_tasks:
            t = to.task
            recent_tasks_data.append({
                'id': str(t.id),
                'title': t.title,
                'priority': t.priority,
                'status': to.status,
                'deadline': str(t.deadline) if t.deadline else None,
                'created_at': t.created_at.isoformat(),
                'assigned_to': to.assigned_to.full_name if to.assigned_to else None,
            })
        
        # ---- Murojaat statistikasi ----
        appeals = TelegramAppeal.objects.filter(assigned_organizations=org)
        total_appeals = appeals.count()
        pending_appeals = appeals.filter(status__in=['pending_review', 'pending_ai']).count()
        approved_appeals = appeals.filter(status='approved').count()
        responded_appeals = appeals.filter(status='responded').count()
        resolved_appeals = appeals.filter(status='resolved').count()
        avg_rating = appeals.filter(rating__isnull=False).aggregate(avg=Avg('rating'))['avg']
        rated_count = appeals.filter(rating__isnull=False).count()

        resolved_items = appeals.filter(closed_at__isnull=False)
        resolution_days = []
        for appeal in resolved_items:
            if appeal.closed_at and appeal.created_at:
                resolution_days.append((appeal.closed_at - appeal.created_at).total_seconds() / 86400)
        avg_resolution_days = round(sum(resolution_days) / len(resolution_days), 1) if resolution_days else 0
        
        # Oxirgi 5 ta murojaat
        recent_appeals = appeals.order_by('-created_at')[:5]
        recent_appeals_data = []
        for a in recent_appeals:
            recent_appeals_data.append({
                'id': a.id,
                'appeal_number': a.appeal_number or f'#{a.id}',
                'text': a.text[:120] + '...' if len(a.text) > 120 else a.text,
                'status': a.status,
                'priority': a.priority,
                'created_at': a.created_at.isoformat(),
                'user_name': a.telegram_user.full_name if a.telegram_user else (a.citizen_name or 'Noma\'lum'),
                'category_name': a.category.name_uz if a.category else None,
            })
        
        # ---- Xodimlar ----
        employees = org.employees.filter(status='FAOL')
        employee_count = employees.count()
        
        # ---- Tashkilot ma'lumotlari ----
        org_data = {
            'id': str(org.id),
            'name': org.name,
            'short_name': org.short_name,
            'sector': org.sector.name if org.sector else None,
            'director_name': org.director_name,
            'address': org.address,
            'phone': org.phone,
        }
        
        return Response({
            'organization': org_data,
            'tasks': {
                'total': total_tasks,
                'new': new_tasks,
                'in_progress': in_progress,
                'in_review': in_review,
                'completed': completed,
                'overdue': overdue,
                'resubmitted': resubmitted,
                'completion_rate': completion_rate,
                'recent': recent_tasks_data,
            },
            'appeals': {
                'total': total_appeals,
                'pending': pending_appeals,
                'approved': approved_appeals,
                'responded': responded_appeals,
                'resolved': resolved_appeals,
                'avg_resolution_days': avg_resolution_days,
                'average_rating': round(avg_rating, 1) if avg_rating is not None else None,
                'rated_count': rated_count,
                'recent': recent_appeals_data,
            },
            'service': {
                'target_review_days': 3,
                'target_response_days': 5,
            },
            'employees': {
                'count': employee_count,
            },
        })


class OrganizationAnalyticsView(views.APIView):
    """
    Organization analytics endpoint.
    
    GET /api/analytics/organizations/
    """
    permission_classes = [IsAuthenticated, IsHokimOrHokimlikMasul]
    
    def get(self, request):
        # 4N+1 tuzatildi — _organization_stats() bitta agregat so'rov ishlatadi.
        stats = _organization_stats()[:5]

        return Response([
            {
                'organization': {
                    'id': row['id'],
                    'name': row['name'],
                    'short_name': row['short_name'],
                },
                'total_tasks': row['total'],
                'completed_tasks': row['completed'],
                'in_progress_tasks': row['in_progress'],
                'overdue_tasks': row['overdue'],
                'completion_rate': row['completion_rate'],
                'rating': row['rating'],
            }
            for row in stats
        ])


class UserAnalyticsView(views.APIView):
    """
    User analytics endpoint.
    
    GET /api/analytics/users/
    """
    permission_classes = [IsAuthenticated, IsHokimOrHokimlikMasul]
    
    def get(self, request):
        # Users by role
        users_by_role = User.objects.filter(
            status='FAOL'
        ).values('role').annotate(count=Count('id'))
        
        # Users by status
        users_by_status = User.objects.values('status').annotate(count=Count('id'))
        
        # Top performers (most completed tasks)
        top_performers = User.objects.filter(
            status='FAOL',
            role__in=['TASHKILOT_RAHBARI', 'TASHKILOT_MASUL']
        ).annotate(
            completed_tasks=Count(
                'assigned_tasks',
                filter=Q(assigned_tasks__status__in=['BAJARILDI', 'NAZORATDAN_YECHILDI'])
            )
        ).order_by('-completed_tasks')[:10]
        
        return Response({
            'users_by_role': list(users_by_role),
            'users_by_status': list(users_by_status),
            'top_performers': [
                {
                    'id': str(u.id),
                    'name': u.full_name,
                    'role': u.role,
                    'organization': u.organization.name if u.organization else None,
                    'completed_tasks': u.completed_tasks,
                }
                for u in top_performers
            ],
        })


class TaskTrendsView(views.APIView):
    """
    Task trends analytics endpoint.
    
    GET /api/analytics/trends/
    """
    permission_classes = [IsAuthenticated, IsHokimOrHokimlikMasul]
    
    def get(self, request):
        days = int(request.query_params.get('days', 30))
        start_date = timezone.now().date() - timedelta(days=days)
        user = request.user
        
        # Tashkilot xodimlari uchun TaskOrganization dan hisoblash
        if user.role in ['TASHKILOT_RAHBARI', 'TASHKILOT_MASUL'] and user.organization:
            from tasks.models import TaskOrganization
            task_orgs = TaskOrganization.objects.filter(organization=user.organization)
            
            # Tasks created per day (tayinlangan sana bo'yicha)
            created_trend = task_orgs.filter(
                assigned_at__date__gte=start_date
            ).annotate(
                date=TruncDate('assigned_at')
            ).values('date').annotate(
                count=Count('id')
            ).order_by('date')
            
            # Tasks completed per day
            completed_trend = task_orgs.filter(
                status='BAJARILDI',
                completed_at__date__gte=start_date
            ).annotate(
                date=TruncDate('completed_at')
            ).values('date').annotate(
                count=Count('id')
            ).order_by('date')
            
            # Status distribution
            status_distribution = task_orgs.values('status').annotate(count=Count('id'))
        else:
            # Admin rollar uchun Task dan hisoblash
            # Tasks created per day
            created_trend = Task.objects.filter(
                created_at__date__gte=start_date
            ).annotate(
                date=TruncDate('created_at')
            ).values('date').annotate(
                count=Count('id')
            ).order_by('date')
            
            # Tasks completed per day
            completed_trend = Task.objects.filter(
                status__in=['BAJARILDI', 'NAZORATDAN_YECHILDI'],
                updated_at__date__gte=start_date
            ).annotate(
                date=TruncDate('updated_at')
            ).values('date').annotate(
                count=Count('id')
            ).order_by('date')
            
            # Tasks by status over time
            status_distribution = Task.objects.values('status').annotate(count=Count('id'))
        
        # Average completion time
        completed_tasks = Task.objects.filter(
            closed_at__isnull=False
        ).annotate(
            completion_time=F('closed_at') - F('created_at')
        )
        
        avg_completion_days = 0
        if completed_tasks.exists():
            total_days = sum(
                (t.completion_time.total_seconds() / 86400) 
                for t in completed_tasks 
                if t.completion_time
            )
            avg_completion_days = total_days / completed_tasks.count()
        
        return Response({
            'created_trend': list(created_trend),
            'completed_trend': list(completed_trend),
            'status_distribution': list(status_distribution),
            'average_completion_days': round(avg_completion_days, 1),
        })


class AnalyticsExportView(views.APIView):
    """
    Analytics export endpoint.
    
    GET /api/analytics/export/?format=xlsx|pdf
    """
    permission_classes = [IsAuthenticated, IsHokimOrHokimlikMasul]

    def get(self, request):
        export_format = request.query_params.get('format', 'xlsx').lower()

        data = self._collect_data()

        if export_format == 'pdf':
            return self._export_pdf(data)

        return self._export_xlsx(data)

    def _collect_data(self):
        today = timezone.now().date()
        thirty_days_ago = today - timedelta(days=30)

        tasks = Task.objects.all()

        # Ilgari bu blok yettita alohida .count() so'rovi yuborardi.
        # Endi bitta aggregate. `pending_approval` ham to'g'rilandi: u
        # YANGI ni sanardi, holbuki tasdiqlash kutayotgan holat BAJARILDI.
        agg = tasks.aggregate(
            total=Count('id'),
            new=Count('id', filter=Q(status='YANGI')),
            in_progress=Count('id', filter=Q(status='IJRODA')),
            awaiting_approval=Count('id', filter=Q(status='BAJARILDI')),
            closed=Count('id', filter=Q(status='NAZORATDAN_YECHILDI')),
            resubmitted=Count('id', filter=Q(status='QAYTA_IJROGA_YUBORILDI')),
            overdue=Count('id', filter=Q(status='MUDDATI_KECH')),
        )
        user_agg = User.objects.aggregate(
            total=Count('id'),
            active=Count('id', filter=Q(status='FAOL')),
        )

        summary = {
            'total_tasks': agg['total'] or 0,
            'new_tasks': agg['new'] or 0,
            'in_progress_tasks': agg['in_progress'] or 0,
            'completed_tasks': (agg['awaiting_approval'] or 0) + (agg['closed'] or 0),
            'pending_approval': agg['awaiting_approval'] or 0,
            'resubmitted': agg['resubmitted'] or 0,
            'overdue_tasks': agg['overdue'] or 0,
            'total_users': user_agg['total'] or 0,
            'active_users': user_agg['active'] or 0,
            'total_organizations': Organization.objects.filter(is_active=True).count(),
        }

        tasks_by_priority = list(
            tasks.values('priority').annotate(count=Count('id')).order_by('priority')
        )

        tasks_trend = list(
            tasks.filter(created_at__date__gte=thirty_days_ago)
            .annotate(date=TruncDate('created_at'))
            .values('date')
            .annotate(created=Count('id'))
            .order_by('date')
        )

        completed_trend = list(
            tasks.filter(closed_at__date__gte=thirty_days_ago)
            .annotate(date=TruncDate('closed_at'))
            .values('date')
            .annotate(completed=Count('id'))
            .order_by('date')
        )

        # Xuddi shu 4N+1 eksport yo'lida ham bor edi — endi umumiy yordamchi.
        org_stats = [
            {
                'name': row['name'],
                'total_tasks': row['total'],
                'completed_tasks': row['completed'],
                'in_progress_tasks': row['in_progress'],
                'overdue_tasks': row['overdue'],
                'completion_rate': row['completion_rate'],
                'rating': row['rating'],
            }
            for row in _organization_stats()
        ]

        users_by_role = list(User.objects.filter(status='FAOL').values('role').annotate(count=Count('id')))
        users_by_status = list(User.objects.values('status').annotate(count=Count('id')))
        top_performers = list(
            User.objects.filter(
                status='FAOL',
                role__in=['TASHKILOT_RAHBARI', 'TASHKILOT_MASUL']
            )
            .annotate(
                completed_tasks=Count(
                    'assigned_tasks',
                    filter=Q(assigned_tasks__status__in=['BAJARILDI', 'NAZORATDAN_YECHILDI'])
                )
            )
            .order_by('-completed_tasks')
            # `full_name` — property, DB maydoni emas: uni values() ga berish
            # FieldError beradi. Haqiqiy maydonlarni olib, ismni Python
            # tomonida yig'amiz (javob shakli o'zgarmaydi).
            .values('last_name', 'first_name', 'middle_name', 'role', 'completed_tasks')[:10]
        )
        top_performers = [
            {
                'full_name': ' '.join(
                    part for part in (row['last_name'], row['first_name'], row['middle_name'])
                    if part
                ),
                'role': row['role'],
                'completed_tasks': row['completed_tasks'],
            }
            for row in top_performers
        ]

        return {
            'summary': summary,
            'tasks_by_priority': tasks_by_priority,
            'tasks_trend': tasks_trend,
            'completed_trend': completed_trend,
            'org_stats': org_stats,
            'users_by_role': users_by_role,
            'users_by_status': users_by_status,
            'top_performers': top_performers,
        }

    def _export_xlsx(self, data: dict):
        wb = Workbook()

        # Summary sheet
        ws = wb.active
        ws.title = 'Summary'
        ws.append(['Metric', 'Value'])
        for key, value in data['summary'].items():
            ws.append([key, value])

        # Tasks by priority
        ws_priority = wb.create_sheet('Tasks Priority')
        ws_priority.append(['Priority', 'Count'])
        for row in data['tasks_by_priority']:
            ws_priority.append([row['priority'], row['count']])
        if len(data['tasks_by_priority']) > 0:
            chart = BarChart()
            chart.title = 'Tasks by Priority'
            data_ref = Reference(ws_priority, min_col=2, min_row=1, max_row=ws_priority.max_row)
            cats_ref = Reference(ws_priority, min_col=1, min_row=2, max_row=ws_priority.max_row)
            chart.add_data(data_ref, titles_from_data=True)
            chart.set_categories(cats_ref)
            ws_priority.add_chart(chart, 'D2')

        # Trends sheet
        ws_trends = wb.create_sheet('Trends')
        ws_trends.append(['Date', 'Created', 'Completed'])
        completed_map = {row['date']: row['completed'] for row in data['completed_trend']}
        for row in data['tasks_trend']:
            ws_trends.append([row['date'], row['created'], completed_map.get(row['date'], 0)])
        if ws_trends.max_row > 2:
            line_chart = LineChart()
            line_chart.title = 'Tasks Trend (Last 30 days)'
            data_ref = Reference(ws_trends, min_col=2, min_row=1, max_col=3, max_row=ws_trends.max_row)
            cats_ref = Reference(ws_trends, min_col=1, min_row=2, max_row=ws_trends.max_row)
            line_chart.add_data(data_ref, titles_from_data=True)
            line_chart.set_categories(cats_ref)
            ws_trends.add_chart(line_chart, 'E2')

        # Organizations sheet
        ws_org = wb.create_sheet('Organizations')
        ws_org.append([
            'Organization', 'Total', 'Completed', 'In Progress', 'Overdue', 'Completion Rate', 'Rating'
        ])
        for row in data['org_stats']:
            ws_org.append([
                row['name'], row['total_tasks'], row['completed_tasks'], row['in_progress_tasks'],
                row['overdue_tasks'], row['completion_rate'], row['rating']
            ])
        if ws_org.max_row > 2:
            bar_chart = BarChart()
            bar_chart.title = 'Organization Ratings'
            data_ref = Reference(ws_org, min_col=7, min_row=1, max_row=ws_org.max_row)
            cats_ref = Reference(ws_org, min_col=1, min_row=2, max_row=ws_org.max_row)
            bar_chart.add_data(data_ref, titles_from_data=True)
            bar_chart.set_categories(cats_ref)
            ws_org.add_chart(bar_chart, 'I2')

        # Users sheet
        ws_users = wb.create_sheet('Users')
        ws_users.append(['Role', 'Count'])
        for row in data['users_by_role']:
            ws_users.append([row['role'], row['count']])
        ws_users.append([])
        ws_users.append(['Status', 'Count'])
        for row in data['users_by_status']:
            ws_users.append([row['status'], row['count']])
        ws_users.append([])
        ws_users.append(['Top Performer', 'Role', 'Completed Tasks'])
        for row in data['top_performers']:
            ws_users.append([row['full_name'], row['role'], row['completed_tasks']])

        output = BytesIO()
        wb.save(output)
        output.seek(0)

        response = HttpResponse(
            output.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename="analytics_export.xlsx"'
        return response

    def _export_pdf(self, data: dict):
        buffer = BytesIO()
        pdf = canvas.Canvas(buffer, pagesize=A4)
        width, height = A4

        pdf.setFont('Helvetica-Bold', 14)
        pdf.drawString(40, height - 40, 'Analytics Report')

        pdf.setFont('Helvetica', 10)
        y = height - 70

        pdf.drawString(40, y, 'Summary')
        y -= 16
        for key, value in data['summary'].items():
            pdf.drawString(50, y, f"{key}: {value}")
            y -= 12
            if y < 80:
                pdf.showPage()
                pdf.setFont('Helvetica', 10)
                y = height - 40

        pdf.showPage()
        pdf.save()

        buffer.seek(0)
        response = HttpResponse(buffer.getvalue(), content_type='application/pdf')
        response['Content-Disposition'] = 'attachment; filename="analytics_report.pdf"'
        return response

class VillageAnalyticsView(views.APIView):
    """
    Qishloqlar kesimidagi agregat — interaktiv xarita shu bilan bo'yaladi.

    GET /api/analytics/villages/
    200 { "<qishloq_kodi>": { "appeals_total": 18, "appeals_open": 4,
                              "appeals_closed": 14 }, ... }

    Kalit — `BotRegion.code`. Xarita geometriyasidagi `code` maydoni ham
    aynan shu (public/geo/xatirchi-villages.paths.json). Agar kodlar mos
    kelmasa, admin panelida `BotRegion.code` ni geometriya kodiga
    tenglashtirish kifoya — migratsiya kerak emas.

    MUHIM CHEKLOV: `Task` modelida qishloq/hudud o'lchovi YO'Q — faqat
    erkin `address` va ixtiyoriy `latitude`/`longitude`. Shuning uchun bu
    endpoint FAQAT murojaatlar bo'yicha raqam qaytaradi. Topshiriqlarni
    ham qishloq kesimida ko'rsatish uchun `Task` ga `region` FK qo'shish
    kerak; u qo'shilmagan holda nol qiymatlarni "ma'lumot" sifatida
    chiqarish xato bo'lardi, shuning uchun ular umuman yuborilmaydi.

    Ikkita agregat so'rov: murojaat Telegram foydalanuvchisi orqali
    (`telegram_user__region`) yoki qo'lda kiritilgan (`citizen_region`)
    bo'lishi mumkin.
    """

    permission_classes = [IsAuthenticated]

    def get(self, request):
        result = {}

        def merge(code, total, closed):
            if not code:
                return
            row = result.setdefault(
                str(code),
                {'appeals_total': 0, 'appeals_open': 0, 'appeals_closed': 0},
            )
            row['appeals_total'] += int(total or 0)
            row['appeals_closed'] += int(closed or 0)
            row['appeals_open'] = row['appeals_total'] - row['appeals_closed']

        telegram_rows = (
            TelegramAppeal.objects
            .filter(telegram_user__region__isnull=False)
            .values('telegram_user__region__code')
            .annotate(
                total=Count('id'),
                closed=Count('id', filter=Q(closed_at__isnull=False)),
            )
        )
        for row in telegram_rows:
            merge(row['telegram_user__region__code'], row['total'], row['closed'])

        manual_rows = (
            TelegramAppeal.objects
            .filter(telegram_user__isnull=True, citizen_region__isnull=False)
            .values('citizen_region__code')
            .annotate(
                total=Count('id'),
                closed=Count('id', filter=Q(closed_at__isnull=False)),
            )
        )
        for row in manual_rows:
            merge(row['citizen_region__code'], row['total'], row['closed'])

        # ------------------------------------------------------------------
        # Hudud passporti: maydon va aholi soni.
        #
        # Xaritada qishloq bosilganda yon panelda to'rt ko'rsatkich chiqadi:
        # maydoni, aholi soni, murojaatlar soni, hal etilganlar soni.
        # Birinchi ikkitasi BotRegion da saqlanadi va TO'LDIRILMAGAN
        # bo'lishi mumkin — bunda maydon umuman yuborilmaydi va frontend
        # "—" ko'rsatadi. Nol yuborish xato bo'lardi.
        #
        # MUHIM: bu yerda barcha hududlar aylanib chiqiladi, shuning uchun
        # murojaati yo'q qishloq ham javobda bo'ladi — xarita ular uchun
        # ham nom va passportni ko'rsatadi.
        for region in BotRegion.objects.filter(is_active=True).only(
            'code', 'name_uz', 'name_ru', 'name_en', 'population', 'area_km2'
        ):
            row = result.setdefault(
                str(region.code),
                {'appeals_total': 0, 'appeals_open': 0, 'appeals_closed': 0},
            )
            row['name'] = region.name_uz
            if region.population is not None:
                row['population'] = int(region.population)
            if region.area_km2 is not None:
                row['area_km2'] = float(region.area_km2)

        return Response(result)
